# -*- coding: utf-8 -*-
"""Tuesday target mini-deck + Excel numbers + ML terms cheat slide."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

OUT = Path(__file__).resolve().parent

# ---------- Excel ----------
wb = Workbook()

# Sheet 1: rule
ws = wb.active
ws.title = "Правило target"
header = PatternFill("solid", fgColor="1F4E79")
header_font = Font(color="FFFFFF", bold=True, size=12)
thin = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

rows = [
    ["Вопрос", "Ответ"],
    ["Что такое bad (target=1)", "Просрочка платежа >= 30 дней (DPD30+)"],
    ["Горизонт", "Первые 6 месяцев после выдачи (issue_date)"],
    ["Как считаем DPD", "paid_date - due_date; если не оплатил: 2025-05-31 - due_date"],
    ["Почему 30 дней", "Мелкая задержка не = плохой клиент; 30 дней — серьёзная просрочка"],
    ["Почему 6 месяцев", "Одинаковое окно для всех; ранний риск; хватает зрелых когорт"],
    ["Отказы", "НЕ в target (6723 шт). Нет договора → нет факта оплаты"],
    ["Незрелые договоры", "ИСКЛЮЧИЛИ (2456 шт), если issue+6м > 2025-05-31"],
    ["Сколько в target", "8821 договор"],
    ["Сколько bad", "2113 (23.95%)"],
    ["Вывод для бизнеса", "Просрочка выросла: ~20% → пик окт 2024 = 36.3%"],
    ["Snapshot данных", "2025-05-31"],
]
for r in rows:
    ws.append(r)
for cell in ws[1]:
    cell.fill = header
    cell.font = header_font
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2):
    for c in row:
        c.border = thin
        c.alignment = Alignment(wrap_text=True, vertical="center")
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 78
ws.row_dimensions[1].height = 22

# Sheet 2: cohorts
ws2 = wb.create_sheet("Bad по когортам")
ws2.append(["Месяц выдачи", "N договоров", "Bad", "Bad %"])
cohorts = [
    ("2023-06", 432, 106, 24.54),
    ("2023-07", 495, 116, 23.43),
    ("2023-08", 497, 107, 21.53),
    ("2023-09", 488, 85, 17.42),
    ("2023-10", 544, 131, 24.08),
    ("2023-11", 461, 92, 19.96),
    ("2023-12", 501, 99, 19.76),
    ("2024-01", 473, 85, 17.97),
    ("2024-02", 484, 96, 19.83),
    ("2024-03", 492, 101, 20.53),
    ("2024-04", 464, 91, 19.61),
    ("2024-05", 543, 119, 21.92),
    ("2024-06", 482, 122, 25.31),
    ("2024-07", 511, 141, 27.59),
    ("2024-08", 493, 144, 29.21),
    ("2024-09", 452, 128, 28.32),
    ("2024-10", 507, 184, 36.29),
    ("2024-11", 502, 166, 33.07),
    ("ИТОГО", 8821, 2113, 23.95),
]
red = PatternFill("solid", fgColor="FCE4EC")
for c in cohorts:
    ws2.append(list(c))
for cell in ws2[1]:
    cell.fill = header
    cell.font = header_font
for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=4):
    for c in row:
        c.border = thin
    # highlight peak months
    if row[0].value in ("2024-10", "2024-11"):
        for c in row:
            c.fill = red
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14
ws2.column_dimensions["C"].width = 10
ws2.column_dimensions["D"].width = 10

# Sheet 3: speech
ws3 = wb.create_sheet("Текст на показ")
ws3["A1"] = "ЗАЧИТАЙ НА ПОКАЗЕ (60–90 сек)"
ws3["A1"].font = Font(bold=True, size=14, color="1F4E79")
speech = (
    "Мы разметили, кто плохой клиент после выдачи кредита.\n\n"
    "Плохой (target = 1) — если в первые 6 месяцев после выдачи была просрочка платежа на 30 дней или больше.\n\n"
    "Почему 30 дней: опоздание на пару дней бывает у всех. 30 дней — уже серьёзно.\n\n"
    "Почему 6 месяцев: всем даём одинаковое окно. Смотрим ранний риск.\n\n"
    "Отказы в таргет не берём: кредит не выдали — платежей нет — лейбла нет.\n\n"
    "Незрелые договоры исключаем: если 6 месяцев ещё не прошло к 31 мая 2025 — исход неизвестен.\n\n"
    "На 8821 договоре средний bad rate 24%. По когортам видно рост: было около 20%, "
    "в октябре 2024 — 36%. Просрочка действительно выросла."
)
ws3["A3"] = speech
ws3["A3"].alignment = Alignment(wrap_text=True, vertical="top")
ws3.column_dimensions["A"].width = 100
ws3.row_dimensions[3].height = 200

# Sheet 4: ML terms for if asked
ws4 = wb.create_sheet("Если спросят ML-термины")
ws4.append(["Термин", "Простыми словами", "Когда нужен"])
ml_terms = [
    ("AUC-ROC", "Насколько модель отличает плохих от хороших (0.5 = монетка, 1.0 = идеал)", "Среда: качество модели"),
    ("Gini", "Почти то же, что AUC: Gini = 2*AUC − 1. Чем выше — лучше ранжирование", "Среда: часто просят банки"),
    ("CatBoost", "Готовая модель «деревья», хорошо ест таблицы и категории без ручной возни", "Среда: чем учим модель"),
    ("PSI", "Насколько фичи/скор „поплыли“ со временем. Большой PSI = мир изменился", "Среда: стабильность train vs test"),
    ("Out-of-time", "Учим на прошлом, проверяем на будущем. НЕ случайный split", "Среда: обязательно"),
    ("Порог (threshold)", "С какого скора одобряем. Выше порог → меньше одобрений, меньше bad", "Среда: бизнес-вывод"),
    ("Leakage", "Подглядывание в будущее (коллекшн после выдачи, dpd_max и т.п.) — нельзя", "Всегда"),
]
for r in ml_terms:
    ws4.append(list(r))
for cell in ws4[1]:
    cell.fill = header
    cell.font = header_font
for row in ws4.iter_rows(min_row=1, max_row=ws4.max_row, max_col=3):
    for c in row:
        c.border = thin
        c.alignment = Alignment(wrap_text=True, vertical="center")
ws4.column_dimensions["A"].width = 16
ws4.column_dimensions["B"].width = 70
ws4.column_dimensions["C"].width = 28

xlsx_path = OUT / "Scoring_Target_Vtornik.xlsx"
wb.save(xlsx_path)

# ---------- PPTX ----------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0x0C, 0x10, 0x17)
CARD = RGBColor(0x15, 0x1C, 0x27)
ACCENT = RGBColor(0x4D, 0x9F, 0xFF)
BAD = RGBColor(0xE8, 0x5D, 0x5D)
OK = RGBColor(0x3E, 0xCF, 0x8E)
WARN = RGBColor(0xE8, 0xB8, 0x4A)
TEXT = RGBColor(0xE8, 0xEE, 0xF4)
MUTED = RGBColor(0x8B, 0x9A, 0xAB)


def set_slide_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_box(slide, left, top, width, height, fill_rgb=CARD):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def set_run(p, text, size=18, bold=False, color=TEXT):
    p.clear()
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Segoe UI"
    return run


def add_text(slide, left, top, width, height, text, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p, text, size=size, bold=bold, color=color)
    return box


def kpi(slide, left, top, title, value, value_color=TEXT):
    add_box(slide, left, top, Inches(2.9), Inches(1.35))
    add_text(slide, left + Inches(0.15), top + Inches(0.15), Inches(2.6), Inches(0.4), title, size=12, color=MUTED)
    add_text(slide, left + Inches(0.15), top + Inches(0.5), Inches(2.6), Inches(0.7), value, size=28, bold=True, color=value_color)


# Slide 1 title
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s)
add_text(s, Inches(0.7), Inches(2.2), Inches(12), Inches(0.4), "ВТОРНИК · ПОКАЗ ТАРГЕТА", size=14, bold=True, color=ACCENT)
add_text(s, Inches(0.7), Inches(2.7), Inches(12), Inches(1), "Кто плохой клиент после выдачи кредита", size=36, bold=True)
add_text(s, Inches(0.7), Inches(3.8), Inches(11), Inches(0.8),
         "Bad = DPD ≥ 30 дней за первые 6 месяцев  ·  snapshot 2025-05-31", size=18, color=MUTED)

# Slide 2 rule
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s)
add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "1. Правило target", size=28, bold=True)
add_box(s, Inches(0.6), Inches(1.1), Inches(12.1), Inches(2.0))
add_text(s, Inches(0.9), Inches(1.3), Inches(11.5), Inches(1.6),
         "target = 1, если за 6 месяцев после выдачи\nбыла просрочка платежа на 30+ дней.\nИначе target = 0.",
         size=26, bold=True)
kpi(s, Inches(0.6), Inches(3.5), "Зрелых договоров", "8 821", ACCENT)
kpi(s, Inches(3.7), Inches(3.5), "Bad rate", "23.95%", BAD)
kpi(s, Inches(6.8), Inches(3.5), "Отказы вне target", "6 723", WARN)
kpi(s, Inches(9.9), Inches(3.5), "Незрелые выкинули", "2 456", WARN)

# Slide 3 why
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s)
add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "2. Почему так", size=28, bold=True)
add_box(s, Inches(0.6), Inches(1.1), Inches(5.9), Inches(5.3))
add_text(s, Inches(0.9), Inches(1.3), Inches(5.3), Inches(0.5), "Почему 30 дней", size=20, bold=True, color=ACCENT)
add_text(s, Inches(0.9), Inches(1.9), Inches(5.3), Inches(3.8),
         "• 2–5 дней опоздания бывает у всех\n"
         "• 30+ дней = серьёзная просрочка\n"
         "• Стандарт риска (DPD30)\n\n"
         "Можно было 60/90 — рост просрочки\nвсё равно виден",
         size=18, color=TEXT)
add_box(s, Inches(6.8), Inches(1.1), Inches(5.9), Inches(5.3))
add_text(s, Inches(7.1), Inches(1.3), Inches(5.3), Inches(0.5), "Почему 6 месяцев", size=20, bold=True, color=ACCENT)
add_text(s, Inches(7.1), Inches(1.9), Inches(5.3), Inches(3.8),
         "• Всем одинаковое окно\n"
         "• Смотрим ранний риск\n"
         "• Хватает зрелых когорт\n  для проверки модели завтра\n\n"
         "Если взять 12 месяцев —\nслишком много договоров\nещё „не дозрели“",
         size=18, color=TEXT)

# Slide 4 rejects & immature
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s)
add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "3. Отказы и незрелые", size=28, bold=True)
add_box(s, Inches(0.6), Inches(1.1), Inches(5.9), Inches(5.3))
add_text(s, Inches(0.9), Inches(1.3), Inches(5.3), Inches(0.5), "Отказы — НЕ в target", size=22, bold=True, color=WARN)
add_text(s, Inches(0.9), Inches(2.1), Inches(5.3), Inches(3.5),
         "6 723 заявки отклонены.\n\n"
         "Договора нет → платежей нет →\nне знаем, стал бы клиент плохим.\n\n"
         "Их смотрим только чтобы понять,\nкого резала старая политика.",
         size=18)
add_box(s, Inches(6.8), Inches(1.1), Inches(5.9), Inches(5.3))
add_text(s, Inches(7.1), Inches(1.3), Inches(5.3), Inches(0.5), "Незрелые — ИСКЛЮЧИЛИ", size=22, bold=True, color=WARN)
add_text(s, Inches(7.1), Inches(2.1), Inches(5.3), Inches(3.5),
         "2 456 договоров:\nissue + 6 месяцев > 31.05.2025\n\n"
         "Окно ещё не закрылось =\nисход неизвестен.\n\n"
         "Иначе смешаем «ещё не успел»\nс «хороший».",
         size=18)

# Slide 5 cohorts
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s)
add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "4. Bad rate по когортам выдачи", size=28, bold=True)
add_box(s, Inches(0.6), Inches(1.0), Inches(12.1), Inches(5.6))
add_text(s, Inches(0.9), Inches(1.2), Inches(11.5), Inches(0.5),
         "Было ~18–24%  →  рост со 2П 2024  →  пик октябрь 2024 = 36.3%",
         size=20, bold=True, color=BAD)
# simple text bar chart using lines
bars = [
    ("2023 H2", 21),
    ("2024 H1", 20),
    ("2024-07", 28),
    ("2024-08", 29),
    ("2024-09", 28),
    ("2024-10", 36),
    ("2024-11", 33),
]
y0 = 2.0
for i, (lab, val) in enumerate(bars):
    top = Inches(y0 + i * 0.55)
    add_text(s, Inches(1.0), top, Inches(1.6), Inches(0.4), lab, size=14, color=MUTED)
    w = Inches(val / 36 * 8.5)
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.8), top + Inches(0.05), w, Inches(0.35))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BAD if val >= 30 else (WARN if val >= 25 else OK)
    shape.line.fill.background()
    add_text(s, Inches(2.8) + w + Inches(0.15), top, Inches(1.2), Inches(0.4), f"{val}%", size=14, bold=True)

# Slide 6 speech
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s)
add_text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), "5. Что сказать (зачитай)", size=28, bold=True)
add_box(s, Inches(0.6), Inches(1.0), Inches(12.1), Inches(5.8))
add_text(s, Inches(0.95), Inches(1.25), Inches(11.4), Inches(5.3),
         "Bad — просрочка 30+ дней в первые 6 месяцев после выдачи.\n\n"
         "Отказы не размечаем: нет договора. Незрелые исключаем: окно не закрылось.\n\n"
         "На 8 821 договоре bad = 24%. По когортам рост: с ~20% до 36% в октябре 2024.\n\n"
         "Вывод: просрочка действительно выросла. Завтра на этом таргете — модель.",
         size=22)

# Slide 7 if asked ML
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s)
add_text(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.45), "Если спросят про Gini / CatBoost / AUC / PSI", size=24, bold=True)
add_text(s, Inches(0.6), Inches(0.75), Inches(12), Inches(0.4),
         "Это про СРЕДУ (модель). Сегодня таргет. Но если спросит — скажи так:",
         size=14, color=MUTED)

terms = [
    ("AUC-ROC", "Умеет ли модель отличать плохих от хороших.\n0.5 = случайно, чем ближе к 1 — тем лучше."),
    ("Gini", "Почти AUC для банков: Gini = 2×AUC − 1.\nВыше Gini → лучше ранжирование."),
    ("CatBoost", "Готовый алгоритм деревьев для таблиц.\nИм обычно и учат скоринг."),
    ("PSI", "Не „поплыли“ ли данные со временем.\nБольшой PSI = на тесте мир другой."),
]
positions = [(0.6, 1.3), (6.9, 1.3), (0.6, 4.2), (6.9, 4.2)]
for (title, body), (l, t) in zip(terms, positions):
    add_box(s, Inches(l), Inches(t), Inches(5.8), Inches(2.5))
    add_text(s, Inches(l + 0.25), Inches(t + 0.2), Inches(5.3), Inches(0.45), title, size=20, bold=True, color=ACCENT)
    add_text(s, Inches(l + 0.25), Inches(t + 0.75), Inches(5.3), Inches(1.5), body, size=15)

pptx_path = OUT / "Scoring_Target_Vtornik.pptx"
prs.save(pptx_path)
print("saved", xlsx_path)
print("saved", pptx_path)
